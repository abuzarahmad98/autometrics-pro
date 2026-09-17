"""
Arrow Editor — Flask Web Application
=====================================
A visual tool for configuring arrow annotations on 2D car model images.
Open in browser, click to place arrow start/end points, type a label,
and save the configuration to arrow_config.json.

Supports per-view arrow configurations (side_view, front_view, top_view)
and dynamically resolves images based on the master car type.

Usage:
    python arrow_editor/app.py
    Then open http://localhost:5000 in your browser.
"""

import os
import sys
import json
import webbrowser
from threading import Timer

from flask import Flask, render_template, jsonify, request, send_from_directory

# ─── Paths ───────────────────────────────────────────────────────
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
IMAGES_DIR = os.path.join(DATA_DIR, "images")
CONFIG_DIR = os.path.join(BASE_DIR, "config")
ARROW_CONFIG_PATH = os.path.join(CONFIG_DIR, "arrow_config.json")
EXCEL_PATH = os.path.join(DATA_DIR, "car_metrics.xlsx")

# ─── View types ──────────────────────────────────────────────────
VIEW_TYPES = ["side_view", "front_view", "top_view"]

# ─── Car type to image mapping (type -> view -> filename) ───────
TYPE_VIEW_IMAGE_MAP = {
    "Boxy Master": {
        "side_view":  "boxy_side_view.png",
        "front_view": "boxy_front_view.png",
        "top_view":   "boxy_top_view.png",
    },
    "Curvy Master": {
        "side_view":  "curvy_side_view.png",
        "front_view": "curvy_front_view.png",
        "top_view":   "curvy_top_view.png",
    },
    "Coupe Master": {
        "side_view":  "coupe_side_view.png",
        "front_view": "coupe_front_view.png",
        "top_view":   "coupe_top_view.png",
    },
    "Benchmark": {
        "side_view":  "side_view.png",
        "front_view": "side_view.png",
        "top_view":   "side_view.png",
    },
}

app = Flask(__name__)


def load_arrow_config():
    """Load existing arrow configuration.

    Auto-migrates old flat format to new per-view format in memory.
    """
    if os.path.exists(ARROW_CONFIG_PATH):
        with open(ARROW_CONFIG_PATH, "r", encoding="utf-8") as f:
            config = json.load(f)

        # Auto-migrate old format entries
        migrated = {}
        for metric_name, metric_data in config.items():
            if "start" in metric_data and "view_type" not in metric_data:
                # Old flat format
                migrated[metric_name] = {
                    "view_type": "side_view",
                    "side_view": {
                        "start": metric_data["start"],
                        "end": metric_data["end"],
                        "label": metric_data.get("label", metric_name),
                        "color": metric_data.get("color", "#E31937"),
                    }
                }
            else:
                migrated[metric_name] = metric_data

        return migrated
    return {}


def save_arrow_config(config):
    """Save arrow configuration to JSON."""
    os.makedirs(CONFIG_DIR, exist_ok=True)
    with open(ARROW_CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(config, f, indent=4, ensure_ascii=False)


def get_master_car_info():
    """Read the master car type from Excel."""
    if not os.path.exists(EXCEL_PATH):
        return None, None

    try:
        import openpyxl
        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        ws_data = wb["Car Data"]
        headers = [cell.value for cell in ws_data[1]]

        for row in ws_data.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))
            car_name = row_dict.get("Car Name", "")
            if not car_name:
                continue
            select = str(row_dict.get("Select", "")).strip().upper()
            if select == "YES":
                car_type = str(row_dict.get("Type", "")).strip()
                wb.close()
                return car_name, car_type

        wb.close()
    except Exception as e:
        print(f"Error reading master car info: {e}")

    return None, None


def get_metrics_from_excel():
    """Read metric names and config from Excel."""
    metrics = []
    metric_config = {}

    if not os.path.exists(EXCEL_PATH):
        return metrics, metric_config

    try:
        import openpyxl
        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

        # Read Car Data headers for metric names
        ws_data = wb["Car Data"]
        headers = [cell.value for cell in ws_data[1]]
        metrics = [h for h in headers if h and h not in ("Select", "Car Name", "Type")]

        # Read Metric Config
        if "Metric Config" in wb.sheetnames:
            ws_config = wb["Metric Config"]
            config_headers = [cell.value for cell in ws_config[1]]
            for row in ws_config.iter_rows(min_row=2, values_only=True):
                row_dict = dict(zip(config_headers, row))
                metric_name = row_dict.get("Metric Name", "")
                if metric_name:
                    metric_config[metric_name] = {
                        "display_name": row_dict.get("Display Name (Alias)", "") or metric_name,
                        "unit": row_dict.get("Unit", "") or "",
                        "image": row_dict.get("Image File", "") or "side_view.png",
                    }

        wb.close()
    except Exception as e:
        print(f"Error reading Excel: {e}")

    return metrics, metric_config


def get_available_images():
    """List all image files in the images directory."""
    images = []
    if os.path.exists(IMAGES_DIR):
        for f in os.listdir(IMAGES_DIR):
            if f.lower().endswith((".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp")):
                images.append(f)
    return sorted(images)


def get_image_for_type_and_view(car_type, view_type):
    """Resolve the correct image filename for a car type + view combination."""
    type_views = TYPE_VIEW_IMAGE_MAP.get(car_type, {})
    image_file = type_views.get(view_type, type_views.get("side_view", "side_view.png"))
    image_path = os.path.join(IMAGES_DIR, image_file)
    if os.path.exists(image_path):
        return image_file
    return "side_view.png"


# ─── Routes ──────────────────────────────────────────────────────

@app.route("/")
def index():
    """Main editor page."""
    return render_template("editor.html")


@app.route("/api/metrics")
def api_metrics():
    """Get list of metrics and their config."""
    metrics, metric_config = get_metrics_from_excel()
    arrow_config = load_arrow_config()

    result = []
    for m in metrics:
        mc = metric_config.get(m, {})
        arrow_data = arrow_config.get(m, {})
        has_arrow = m in arrow_config
        view_type = arrow_data.get("view_type", "side_view") if has_arrow else "side_view"

        # Determine which views have arrow configs
        configured_views = []
        for vt in VIEW_TYPES:
            if vt in arrow_data:
                configured_views.append(vt)

        result.append({
            "name": m,
            "display_name": mc.get("display_name", m),
            "unit": mc.get("unit", ""),
            "image": mc.get("image", "side_view.png"),
            "configured": has_arrow,
            "view_type": view_type,
            "configured_views": configured_views,
        })

    return jsonify(result)


@app.route("/api/master-car")
def api_master_car():
    """Get the current master car name and type."""
    car_name, car_type = get_master_car_info()
    return jsonify({
        "name": car_name,
        "type": car_type,
    })


@app.route("/api/view-types")
def api_view_types():
    """Get available view types."""
    return jsonify(VIEW_TYPES)


@app.route("/api/resolve-image")
def api_resolve_image():
    """Resolve the image filename for a given car type and view type."""
    car_type = request.args.get("car_type", "")
    view_type = request.args.get("view_type", "side_view")
    image_file = get_image_for_type_and_view(car_type, view_type)
    return jsonify({"image": image_file})


@app.route("/api/images")
def api_images():
    """Get list of available images."""
    return jsonify(get_available_images())


@app.route("/api/arrow/<path:metric_name>")
def api_get_arrow(metric_name):
    """Get arrow config for a specific metric (all views)."""
    config = load_arrow_config()
    if metric_name in config:
        return jsonify(config[metric_name])
    return jsonify(None)


@app.route("/api/arrow/<path:metric_name>", methods=["POST"])
def api_save_arrow(metric_name):
    """Save arrow config for a specific metric.

    Expects JSON body with:
        view_type: the active view type for PPT generation
        view: which view this arrow data is for
        start: [x, y] percentages
        end: [x, y] percentages
        label: arrow label text
        color: arrow color hex
    """
    data = request.get_json()
    config = load_arrow_config()

    view = data.get("view", "side_view")
    view_type_active = data.get("view_type", view)

    # Initialize metric entry if needed
    if metric_name not in config:
        config[metric_name] = {}

    # Set the active view type (which view the PPT will use)
    config[metric_name]["view_type"] = view_type_active

    # Save arrow data under the specific view
    config[metric_name][view] = {
        "start": data["start"],
        "end": data["end"],
        "label": data.get("label", metric_name),
        "color": data.get("color", "#E31937"),
    }

    save_arrow_config(config)
    return jsonify({"status": "saved"})


@app.route("/api/arrow/<path:metric_name>", methods=["DELETE"])
def api_delete_arrow(metric_name):
    """Delete arrow config for a specific metric.

    If 'view' query param is provided, only delete that view's config.
    Otherwise delete the entire metric entry.
    """
    view = request.args.get("view")
    config = load_arrow_config()

    if metric_name in config:
        if view and view in config[metric_name]:
            # Delete just one view
            del config[metric_name][view]
            # If no views remain (only view_type key left), remove entirely
            remaining_views = [k for k in config[metric_name] if k in VIEW_TYPES]
            if not remaining_views:
                del config[metric_name]
        else:
            # Delete entire metric
            del config[metric_name]
        save_arrow_config(config)

    return jsonify({"status": "deleted"})


@app.route("/images/<path:filename>")
def serve_image(filename):
    """Serve images from the data/images directory."""
    return send_from_directory(IMAGES_DIR, filename)


def open_browser():
    """Open the browser after a short delay."""
    webbrowser.open("http://localhost:5000")


if __name__ == "__main__":
    print("=" * 50)
    print("  Arrow Editor - Starting...")
    print("  Opening browser at http://localhost:5000")
    print("  Press Ctrl+C to stop")
    print("=" * 50)

    # Open browser after 1.5 seconds
    Timer(1.5, open_browser).start()

    app.run(debug=False, port=5000, host="127.0.0.1")
