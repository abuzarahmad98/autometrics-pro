"""
Mahindra Car Metrics Comparison PPT Generator
==============================================
Reads car data from Excel, generates annotated images with arrows,
bar charts, and difference tables for each metric, and assembles
them into a professional PowerPoint presentation.

Supports master car types (Boxy, Curvy, Coupe) and benchmark cars.
The car marked 'YES' in Select is the primary master (amber yellow bar),
cars marked 'Y' are comparison cars (gray bars).

Usage:
    python generate_ppt.py

Prerequisites:
    - data/car_metrics.xlsx must exist with "Car Data" and "Metric Config" sheets
    - data/images/ must contain the referenced image files
    - config/arrow_config.json should exist (use the Arrow Editor to create it)
"""

import os
import sys
import json
import math
import shutil
from datetime import datetime

import openpyxl
import matplotlib
matplotlib.use("Agg")  # Non-interactive backend
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ─── Paths ───────────────────────────────────────────────────────
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
IMAGES_DIR = os.path.join(DATA_DIR, "images")
CONFIG_DIR = os.path.join(BASE_DIR, "config")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
TEMP_DIR = os.path.join(BASE_DIR, "temp")

EXCEL_PATH = os.path.join(DATA_DIR, "car_metrics.xlsx")
ARROW_CONFIG_PATH = os.path.join(CONFIG_DIR, "arrow_config.json")

# ─── Slide dimensions (16:9) ────────────────────────────────────
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# ─── Bar chart colors ───────────────────────────────────────────
MASTER_BAR_COLOR = "#F5A623"   # Amber yellow for YES car
BENCHMARK_BAR_COLOR = "#A5A5A5"  # Gray for Y cars

# ─── Car type to image mapping (type -> view -> filename) ───────
VIEW_TYPES = ["side_view", "front_view", "top_view"]

TYPE_VIEW_IMAGE_MAP = {
    "Boxy Master": {
        "side_view":  "boxy_side_view.png",
        "front_view": "boxy_front_view.png",
        "top_view":   "boxy_top_view.png",
    },
    "Curvy Master": {
        "side_view":  "curvy_side_view.png",
        "front_view": "curvy_front_view.png",
        "top_view":   "curvy_top_view.png",
    },
    "Coupe Master": {
        "side_view":  "coupe_side_view.png",
        "front_view": "coupe_front_view.png",
        "top_view":   "coupe_top_view.png",
    },
    "Benchmark": {
        "side_view":  "side_view.png",
        "front_view": "side_view.png",
        "top_view":   "side_view.png",
    },
}

# ─── Theme colors ───────────────────────────────────────────────
TITLE_BAR_COLOR = RGBColor(0x00, 0x96, 0x96)   # Teal (similar to reference)
TITLE_TEXT_COLOR = RGBColor(0xFF, 0xFF, 0xFF)    # White
SLIDE_BG_COLOR = RGBColor(0xF5, 0xF5, 0xF5)     # Light grey
BODY_TEXT_COLOR = RGBColor(0x33, 0x33, 0x33)     # Dark grey
TABLE_HEADER_COLOR = RGBColor(0x2F, 0x54, 0x96)  # Dark blue
TABLE_ALT_COLOR = RGBColor(0xE8, 0xF0, 0xFE)     # Light blue tint

# ─── Type group accent colors (for title slide) ────────────────
TYPE_ACCENT_COLORS = {
    "Boxy Master":  RGBColor(0xE6, 0x8A, 0x00),  # Warm orange
    "Curvy Master": RGBColor(0x2F, 0x54, 0x96),  # Deep blue
    "Coupe Master": RGBColor(0x38, 0x8E, 0x3C),  # Green
    "Benchmark":    RGBColor(0x75, 0x75, 0x75),   # Gray
}


def hex_to_rgb(hex_color):
    """Convert hex color string to RGB tuple."""
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))


def read_excel_data():
    """Read car data and metric config from Excel.

    Returns:
        master_car: name of the YES car (or None)
        master_car_type: type of the YES car (e.g. 'Boxy Master')
        selected_cars: list of car names marked Y
        all_cars_data: dict of car_name -> row_dict for all cars
        car_types: dict of car_name -> type string
        type_groups: dict of type -> list of car names
        metric_names: list of metric column names
        metric_config: dict of metric configs from Metric Config sheet
    """
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    # ── Read Car Data sheet ──
    ws_data = wb["Car Data"]
    headers = [cell.value for cell in ws_data[1]]

    master_car = None
    master_car_type = None
    selected_cars = []
    all_cars_data = {}
    car_types = {}
    type_groups = {}

    for row in ws_data.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        car_name = row_dict.get("Car Name", "")
        if not car_name:
            continue  # Skip empty separator rows

        car_type = str(row_dict.get("Type", "")).strip()
        select = str(row_dict.get("Select", "")).strip().upper()

        all_cars_data[car_name] = row_dict
        car_types[car_name] = car_type

        # Track type groups
        if car_type not in type_groups:
            type_groups[car_type] = []
        type_groups[car_type].append(car_name)

        if select == "YES":
            master_car = car_name
            master_car_type = car_type
        elif select == "Y":
            selected_cars.append(car_name)

    # Get metric column names (everything except Select, Type, and Car Name)
    metric_names = [h for h in headers if h and h not in ("Select", "Type", "Car Name")]

    # ── Read Metric Config sheet ──
    metric_config = {}
    if "Metric Config" in wb.sheetnames:
        ws_config = wb["Metric Config"]
        config_headers = [cell.value for cell in ws_config[1]]
        for row in ws_config.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(config_headers, row))
            metric_name = row_dict.get("Metric Name", "")
            if metric_name:
                metric_config[metric_name] = {
                    "display_name": row_dict.get("Display Name (Alias)", "") or metric_name,
                    "unit": row_dict.get("Unit", "") or "",
                    "image": row_dict.get("Image File", "") or "side_view.png",
                }

    wb.close()
    return (master_car, master_car_type, selected_cars, all_cars_data,
            car_types, type_groups, metric_names, metric_config)


def load_arrow_config():
    """Load arrow configuration from JSON.

    Handles both old format (flat) and new format (per-view).
    Auto-migrates old format entries to new per-view format in memory.
    """
    if not os.path.exists(ARROW_CONFIG_PATH):
        print("[WARNING] arrow_config.json not found. Slides will not have arrow annotations.")
        print("          Run the Arrow Editor (run_arrow_editor.bat) to configure arrows.")
        return {}
    with open(ARROW_CONFIG_PATH, "r", encoding="utf-8") as f:
        config = json.load(f)

    # Auto-migrate old format entries
    migrated = {}
    for metric_name, metric_data in config.items():
        if "start" in metric_data and "view_type" not in metric_data:
            # Old flat format — convert to new per-view format
            migrated[metric_name] = {
                "view_type": "side_view",
                "side_view": {
                    "start": metric_data["start"],
                    "end": metric_data["end"],
                    "label": metric_data.get("label", metric_name),
                    "color": metric_data.get("color", "#E31937"),
                }
            }
        else:
            migrated[metric_name] = metric_data

    return migrated


def get_image_for_car_type(car_type, view_type="side_view"):
    """Get the appropriate image file name based on car type and view type.

    Args:
        car_type: The car type string (e.g. 'Boxy Master', 'Curvy Master')
        view_type: The view type ('side_view', 'front_view', 'top_view')

    Returns:
        Image filename string.
    """
    type_views = TYPE_VIEW_IMAGE_MAP.get(car_type, {})
    image_file = type_views.get(view_type, type_views.get("side_view", "side_view.png"))
    image_path = os.path.join(IMAGES_DIR, image_file)
    if os.path.exists(image_path):
        return image_file
    # Fallback to side_view for this type, then generic
    fallback = type_views.get("side_view", "side_view.png")
    if os.path.exists(os.path.join(IMAGES_DIR, fallback)):
        return fallback
    return "side_view.png"


def draw_arrowhead(draw, tip_x, tip_y, angle, size=20, color="red"):
    """Draw an arrowhead at the specified position and angle."""
    left_angle = angle + math.pi * 5 / 6
    right_angle = angle - math.pi * 5 / 6
    left_x = tip_x + size * math.cos(left_angle)
    left_y = tip_y + size * math.sin(left_angle)
    right_x = tip_x + size * math.cos(right_angle)
    right_y = tip_y + size * math.sin(right_angle)
    draw.polygon([(tip_x, tip_y), (left_x, left_y), (right_x, right_y)], fill=color)


def generate_annotated_image(metric_name, arrow_config, metric_config, master_car_type, temp_path):
    """Generate a 2D model image with arrow annotation burned in.

    Uses the image matching the master car's type and the view type
    configured for this metric in the arrow config.
    """
    # Determine view type from arrow config
    metric_arrow = arrow_config.get(metric_name, {})
    view_type = metric_arrow.get("view_type", "side_view")

    # Determine which image to use based on master car type + view type
    image_file = get_image_for_car_type(master_car_type, view_type)
    image_path = os.path.join(IMAGES_DIR, image_file)

    if not os.path.exists(image_path):
        print(f"  [WARNING] Image not found: {image_path}. Skipping annotation.")
        return None

    # Open and convert image
    img = Image.open(image_path).convert("RGBA")
    width, height = img.size

    # Create overlay for the arrow
    overlay = Image.new("RGBA", img.size, (255, 255, 255, 0))
    draw = ImageDraw.Draw(overlay)

    # Get arrow config for this metric (per-view)
    metric_arrow = arrow_config.get(metric_name, {})
    view_type = metric_arrow.get("view_type", "side_view")
    arrow = metric_arrow.get(view_type)
    if arrow:
        start_pct = arrow["start"]
        end_pct = arrow["end"]
        label = arrow.get("label", metric_name)
        color_hex = arrow.get("color", "#E31937")
        color_rgb = hex_to_rgb(color_hex)

        # Convert percentage coordinates to pixel coordinates
        x1 = int(start_pct[0] * width)
        y1 = int(start_pct[1] * height)
        x2 = int(end_pct[0] * width)
        y2 = int(end_pct[1] * height)

        # Draw the arrow line
        line_width = max(3, int(height * 0.005))
        draw.line([(x1, y1), (x2, y2)], fill=color_rgb, width=line_width)

        # Draw arrowheads at both ends (double-headed arrow)
        arrow_size = max(15, int(height * 0.025))
        angle = math.atan2(y2 - y1, x2 - x1)

        # Arrowhead at end point
        draw_arrowhead(draw, x2, y2, angle, arrow_size, color_rgb)
        # Arrowhead at start point (reversed)
        draw_arrowhead(draw, x1, y1, angle + math.pi, arrow_size, color_rgb)

        # Draw label
        font_size = max(16, int(height * 0.03))
        try:
            font = ImageFont.truetype("arial.ttf", font_size)
        except (IOError, OSError):
            try:
                font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", font_size)
            except (IOError, OSError):
                font = ImageFont.load_default()

        # Position label near the midpoint of the arrow
        mid_x = (x1 + x2) // 2
        mid_y = (y1 + y2) // 2

        # Get text bounding box
        bbox = draw.textbbox((0, 0), label, font=font)
        text_width = bbox[2] - bbox[0]
        text_height = bbox[3] - bbox[1]

        # Offset label slightly from the arrow line
        label_x = mid_x - text_width // 2
        label_y = mid_y - text_height - 8

        # Draw label background
        padding = 4
        draw.rectangle(
            [label_x - padding, label_y - padding,
             label_x + text_width + padding, label_y + text_height + padding],
            fill=(255, 255, 255, 220)
        )
        draw.text((label_x, label_y), label, fill=color_rgb, font=font)

    # Composite and save
    result = Image.alpha_composite(img, overlay)
    result = result.convert("RGB")
    result.save(temp_path, quality=95)
    return temp_path


def generate_bar_chart(metric_name, master_car, selected_cars, all_cars_data,
                       metric_config, temp_path):
    """Generate a bar chart with master car (amber) and comparison cars (gray).

    The master (YES) car's bar comes first, followed by a 2-bar gap,
    then all Y comparison cars sorted by value descending.
    """
    # Get display name and unit
    config = metric_config.get(metric_name, {})
    display_name = config.get("display_name", metric_name)
    unit = config.get("unit", "")

    # Get master car value
    master_val = 0.0
    if master_car:
        val = all_cars_data.get(master_car, {}).get(metric_name)
        if val is not None:
            try:
                master_val = float(val)
            except (ValueError, TypeError):
                master_val = 0.0

    # Get comparison car values
    comp_values = []
    for car in selected_cars:
        val = all_cars_data.get(car, {}).get(metric_name)
        if val is not None:
            try:
                val = float(val)
            except (ValueError, TypeError):
                val = 0.0
        else:
            val = 0.0
        comp_values.append((car, val))

    # Sort comparison cars by value descending
    comp_values.sort(key=lambda x: x[1], reverse=True)

    # Build bar data: master + gap + comparisons
    bar_names = []
    bar_values = []
    bar_colors = []

    if master_car:
        bar_names.append(master_car)
        bar_values.append(master_val)
        bar_colors.append(MASTER_BAR_COLOR)

        # 2-bar gap (invisible bars)
        bar_names.append("")
        bar_values.append(0)
        bar_colors.append("none")
        bar_names.append("")
        bar_values.append(0)
        bar_colors.append("none")

    for car, val in comp_values:
        bar_names.append(car)
        bar_values.append(val)
        bar_colors.append(BENCHMARK_BAR_COLOR)

    # Create figure
    fig_width = max(8, 2 + len(bar_names) * 0.8)
    fig, ax = plt.subplots(figsize=(fig_width, 3.5))
    fig.patch.set_facecolor("#F5F5F5")
    ax.set_facecolor("#F5F5F5")

    x_positions = range(len(bar_names))

    bars = ax.bar(x_positions, bar_values, color=bar_colors, width=0.6,
                  edgecolor="white", linewidth=1.2)

    # Make gap bars invisible
    for i, bar_obj in enumerate(bars):
        if bar_colors[i] == "none":
            bar_obj.set_visible(False)

    # Get all real (non-gap) values for scaling
    real_values = [v for v, c in zip(bar_values, bar_colors) if c != "none" and v > 0]
    if not real_values:
        real_values = [0]

    # Add value labels on top of visible bars
    for i, (bar_obj, val) in enumerate(zip(bars, bar_values)):
        if bar_colors[i] == "none":
            continue
        # Format: integer if whole number, else 1 decimal
        if val == int(val):
            label = str(int(val))
        else:
            label = f"{val:.1f}"
        ax.text(
            bar_obj.get_x() + bar_obj.get_width() / 2,
            bar_obj.get_height() + (max(real_values) - min(real_values)) * 0.02,
            label,
            ha="center", va="bottom",
            fontsize=10, fontweight="bold", color="#333333"
        )

    # Title
    title_text = display_name
    if unit:
        title_text += f" ({unit})"
    ax.set_title(title_text, fontsize=13, fontweight="bold", color="#009696", pad=12)

    # X-axis labels
    ax.set_xticks(list(x_positions))
    ax.set_xticklabels(bar_names, fontsize=9, rotation=0)

    # Y-axis: start near minimum to emphasize differences
    if real_values and max(real_values) > 0:
        val_range = max(real_values) - min(real_values)
        if val_range > 0:
            y_min = min(real_values) - val_range * 0.3
            y_max = max(real_values) + val_range * 0.15
            if y_min < 0 and min(real_values) >= 0:
                y_min = 0
            ax.set_ylim(y_min, y_max)
        else:
            # All values are the same
            ax.set_ylim(min(real_values) * 0.9, max(real_values) * 1.1)

    # Styling
    ax.tick_params(axis="x", labelsize=9, rotation=0)
    ax.tick_params(axis="y", labelsize=9)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#CCCCCC")
    ax.spines["bottom"].set_color("#CCCCCC")
    ax.yaxis.set_major_locator(ticker.MaxNLocator(nbins=6))
    ax.grid(axis="y", linestyle="--", alpha=0.3)

    plt.tight_layout()
    fig.savefig(temp_path, dpi=180, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return temp_path


def build_difference_data(metric_name, master_car, selected_cars, all_cars_data):
    """Build a sorted difference table for the metric.

    Includes both the master car and selected comparison cars.
    """
    all_compare_cars = []
    if master_car:
        all_compare_cars.append(master_car)
    all_compare_cars.extend(selected_cars)

    car_values = []
    for car in all_compare_cars:
        val = all_cars_data.get(car, {}).get(metric_name)
        if val is not None:
            try:
                val = float(val)
            except (ValueError, TypeError):
                val = 0.0
        else:
            val = 0.0
        car_values.append((car, val))

    # Sort ascending for difference table
    car_values.sort(key=lambda x: x[1])

    diff_data = []
    for i, (car, val) in enumerate(car_values):
        if i == 0:
            diff = "-"
        else:
            d = val - car_values[i - 1][1]
            if d == int(d):
                diff = str(int(d))
            else:
                diff = f"{d:.1f}"
        # Format value
        if val == int(val):
            val_str = str(int(val))
        else:
            val_str = f"{val:.1f}"
        diff_data.append((i + 1, car, val_str, diff))

    return diff_data


def add_title_bar(slide, title_text, subtitle_text=""):
    """Add a styled title bar to the top of a slide."""
    # Title bar background
    title_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        SLIDE_WIDTH, Inches(0.7)
    )
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = TITLE_BAR_COLOR
    title_shape.line.fill.background()

    # Title text
    tf = title_shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TITLE_TEXT_COLOR
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.05)

    # Subtitle on the right side
    if subtitle_text:
        sub_shape = slide.shapes.add_textbox(
            Inches(9.5), Inches(0.1),
            Inches(3.5), Inches(0.5)
        )
        stf = sub_shape.text_frame
        sp = stf.paragraphs[0]
        sp.text = subtitle_text
        sp.font.size = Pt(20)
        sp.font.bold = True
        sp.font.color.rgb = TITLE_TEXT_COLOR
        sp.alignment = PP_ALIGN.RIGHT


def add_difference_table(slide, diff_data, unit, left, top, width, height, master_car=None):
    """Add a difference/ranking table to the slide."""
    rows = len(diff_data) + 1  # +1 for header
    cols = 4  # Rank, Car, Value, Delta

    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    # Set column widths
    table.columns[0].width = Inches(0.5)   # Rank
    table.columns[1].width = Inches(1.2)   # Car Name
    table.columns[2].width = Inches(0.9)   # Value
    table.columns[3].width = Inches(0.7)   # Delta

    # Header row
    header_texts = ["#", "Car", f"Value", "Delta"]
    for col_idx, text in enumerate(header_texts):
        cell = table.cell(0, col_idx)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_HEADER_COLOR
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = PP_ALIGN.CENTER
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Data rows
    for row_idx, (rank, car, val, diff) in enumerate(diff_data, 1):
        row_data = [str(rank), car, val, diff]
        is_master = (car == master_car)

        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(8)
            p.alignment = PP_ALIGN.CENTER
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            if is_master:
                # Highlight master car row with amber tint
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xF0, 0xCC)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0xE6, 0x8A, 0x00)
            elif row_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_ALT_COLOR
                p.font.color.rgb = BODY_TEXT_COLOR
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                p.font.color.rgb = BODY_TEXT_COLOR


def add_bottom_bar(slide, slide_number=None):
    """Add a black line at the bottom of the slide with optional slide number."""
    # Thin black line spanning full width
    line_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(7.15),
        SLIDE_WIDTH, Inches(0.02)
    )
    line_shape.fill.solid()
    line_shape.fill.fore_color.rgb = RGBColor(0x00, 0x00, 0x00)
    line_shape.line.fill.background()

    # Slide number below the line, right-aligned
    if slide_number is not None:
        num_box = slide.shapes.add_textbox(
            Inches(11.5), Inches(7.18),
            Inches(1.5), Inches(0.3)
        )
        ntf = num_box.text_frame
        np_text = ntf.paragraphs[0]
        np_text.text = str(slide_number)
        np_text.font.size = Pt(16)
        np_text.font.bold = True
        np_text.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        np_text.alignment = PP_ALIGN.RIGHT


def create_title_slide(prs, master_car, master_car_type, selected_cars,
                       type_groups, car_types, metric_count):
    """Create a redesigned cover/title slide showing master/benchmark groupings."""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Top decorative bar
    top_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        SLIDE_WIDTH, Inches(1.2)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = TITLE_BAR_COLOR
    top_bar.line.fill.background()

    # Title text
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.15),
        Inches(12), Inches(0.9)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Car Metrics Comparison"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TITLE_TEXT_COLOR
    p.alignment = PP_ALIGN.LEFT

    # Subtitle - date
    date_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(1.4),
        Inches(12), Inches(0.5)
    )
    dtf = date_box.text_frame
    dp = dtf.paragraphs[0]
    dp.text = datetime.now().strftime("Generated on %B %d, %Y at %I:%M %p")
    dp.font.size = Pt(14)
    dp.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # ── Master Car Highlight ──
    if master_car:
        master_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.5), Inches(2.1),
            Inches(5.5), Inches(0.9)
        )
        master_box.fill.solid()
        master_box.fill.fore_color.rgb = RGBColor(0xFF, 0xF0, 0xCC)
        master_box.line.color.rgb = RGBColor(0xF5, 0xA6, 0x23)
        master_box.line.width = Pt(2)

        mtf = master_box.text_frame
        mtf.word_wrap = True
        mp = mtf.paragraphs[0]
        mp.text = f"★  Master Car: {master_car}"
        mp.font.size = Pt(18)
        mp.font.bold = True
        mp.font.color.rgb = RGBColor(0xE6, 0x8A, 0x00)
        mp.alignment = PP_ALIGN.LEFT
        mtf.margin_left = Inches(0.2)

        # Type label
        mp2 = mtf.add_paragraph()
        mp2.text = f"     Type: {master_car_type or 'N/A'}"
        mp2.font.size = Pt(13)
        mp2.font.color.rgb = RGBColor(0x99, 0x66, 0x00)

    # ── Summary info ──
    total_cars = (1 if master_car else 0) + len(selected_cars)
    info_box = slide.shapes.add_textbox(
        Inches(7.0), Inches(2.1),
        Inches(5.5), Inches(0.9)
    )
    itf = info_box.text_frame
    ip = itf.paragraphs[0]
    ip.text = f"Comparing {total_cars} cars across {metric_count} metrics"
    ip.font.size = Pt(15)
    ip.font.bold = True
    ip.font.color.rgb = BODY_TEXT_COLOR

    ip2 = itf.add_paragraph()
    ip2.text = "Advance Vehicle Architecture"
    ip2.font.size = Pt(13)
    ip2.font.color.rgb = TITLE_BAR_COLOR
    ip2.font.bold = True

    # ── Type Groups Section ──
    y_offset = 3.3
    group_order = ["Boxy Master", "Curvy Master", "Coupe Master", "Benchmark"]

    for group_type in group_order:
        cars_in_group = type_groups.get(group_type, [])
        if not cars_in_group:
            continue

        accent_color = TYPE_ACCENT_COLORS.get(group_type, RGBColor(0x66, 0x66, 0x66))

        # Type header bar
        type_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.5), Inches(y_offset),
            Inches(12.0), Inches(0.35)
        )
        type_bar.fill.solid()
        type_bar.fill.fore_color.rgb = accent_color
        type_bar.line.fill.background()

        ttf = type_bar.text_frame
        tp = ttf.paragraphs[0]
        tp.text = f"  {group_type}"
        tp.font.size = Pt(12)
        tp.font.bold = True
        tp.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tp.alignment = PP_ALIGN.LEFT
        ttf.margin_left = Inches(0.15)
        ttf.margin_top = Inches(0.02)

        y_offset += 0.4

        # Car names in this group
        cars_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(y_offset),
            Inches(12), Inches(0.35)
        )
        ctf = cars_box.text_frame
        cp = ctf.paragraphs[0]

        car_labels = []
        for car in cars_in_group:
            if car == master_car:
                car_labels.append(f"★ {car} (YES)")
            elif car in selected_cars:
                car_labels.append(f"{car} (Y)")
            else:
                car_labels.append(car)

        cp.text = "    ".join(car_labels)
        cp.font.size = Pt(12)
        cp.font.color.rgb = BODY_TEXT_COLOR
        cp.alignment = PP_ALIGN.LEFT

        y_offset += 0.5

    # Bottom bar
    add_bottom_bar(slide)

    return slide


def create_metric_slide(prs, metric_name, master_car, master_car_type,
                        selected_cars, all_cars_data, metric_config,
                        arrow_config, slide_number):
    """Create a single metric comparison slide."""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = SLIDE_BG_COLOR

    # Get metric display info
    config = metric_config.get(metric_name, {})
    display_name = config.get("display_name", metric_name)
    unit = config.get("unit", "")

    # ── Title bar ──
    title_text = display_name
    subtitle = "Advance Vehicle Architecture"
    add_title_bar(slide, title_text, subtitle)

    # ── 1. Annotated Image (top-left area) ──
    img_temp_path = os.path.join(TEMP_DIR, f"annotated_{slide_number}.png")
    annotated_path = generate_annotated_image(
        metric_name, arrow_config, metric_config, master_car_type, img_temp_path
    )
    if annotated_path and os.path.exists(annotated_path):
        # Image placement: top-left, below title bar
        slide.shapes.add_picture(
            annotated_path,
            Inches(0.3), Inches(0.85),
            Inches(6.0), Inches(3.0)
        )
    else:
        # Placeholder text if no image
        placeholder = slide.shapes.add_textbox(
            Inches(0.3), Inches(0.85),
            Inches(6.0), Inches(3.0)
        )
        ptf = placeholder.text_frame
        pp = ptf.paragraphs[0]
        pp.text = "[2D Model Image - Configure in Arrow Editor]"
        pp.font.size = Pt(14)
        pp.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        pp.alignment = PP_ALIGN.CENTER

    # ── 2. Bar Chart (bottom-left area) ──
    chart_temp_path = os.path.join(TEMP_DIR, f"chart_{slide_number}.png")
    generate_bar_chart(
        metric_name, master_car, selected_cars, all_cars_data,
        metric_config, chart_temp_path
    )
    slide.shapes.add_picture(
        chart_temp_path,
        Inches(0.2), Inches(4.0),
        Inches(6.5), Inches(3.0)
    )

    # ── 3. Difference Table (right side) ──
    diff_data = build_difference_data(metric_name, master_car, selected_cars, all_cars_data)

    # Table title
    table_title = slide.shapes.add_textbox(
        Inches(7.0), Inches(0.85),
        Inches(3.3), Inches(0.35)
    )
    ttf = table_title.text_frame
    tp = ttf.paragraphs[0]
    tp.text = "Ranking (Ascending)"
    tp.font.size = Pt(11)
    tp.font.bold = True
    tp.font.color.rgb = TABLE_HEADER_COLOR
    tp.alignment = PP_ALIGN.CENTER

    # Calculate table height based on number of rows
    row_height = 0.32
    table_height = (len(diff_data) + 1) * row_height
    max_table_height = 5.5

    add_difference_table(
        slide, diff_data, unit,
        left=Inches(7.0),
        top=Inches(1.25),
        width=Inches(3.3),
        height=Inches(min(table_height, max_table_height)),
        master_car=master_car
    )

    # ── Bottom bar with slide number ──
    add_bottom_bar(slide, slide_number=slide_number)

    return slide


def main():
    """Main entry point for PPT generation."""
    print("=" * 60)
    print("  Mahindra Car Metrics Comparison PPT Generator")
    print("=" * 60)

    # Ensure output and temp directories exist
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs(TEMP_DIR, exist_ok=True)

    # ── Step 1: Read data ──
    print("\n[1/5] Reading Excel data...")
    if not os.path.exists(EXCEL_PATH):
        print(f"  ERROR: Excel file not found at {EXCEL_PATH}")
        print("  Please create the Excel file with car data first.")
        input("\nPress Enter to exit...")
        sys.exit(1)

    (master_car, master_car_type, selected_cars, all_cars_data,
     car_types, type_groups, metric_names, metric_config) = read_excel_data()

    if not master_car and not selected_cars:
        print("  ERROR: No cars selected!")
        print("  Open data/car_metrics.xlsx and put 'YES' for the master car")
        print("  and 'Y' for comparison cars in the Select column.")
        input("\nPress Enter to exit...")
        sys.exit(1)

    if master_car:
        print(f"  Master car (YES): {master_car} [{master_car_type}]")
    else:
        print("  WARNING: No master car (YES) found. Charts will only show Y cars.")

    print(f"  Comparison cars (Y): {', '.join(selected_cars) if selected_cars else 'None'}")
    print(f"  Metrics found: {len(metric_names)}")
    print(f"  Type groups: {', '.join(f'{k} ({len(v)})' for k, v in type_groups.items())}")

    # ── Step 2: Load arrow config ──
    print("\n[2/5] Loading arrow configuration...")
    arrow_config = load_arrow_config()
    configured = sum(1 for m in metric_names if m in arrow_config)
    print(f"  Arrows configured: {configured}/{len(metric_names)} metrics")

    # ── Step 3: Determine images to use ──
    print("\n[3/5] Selecting car type images...")
    if master_car_type:
        for vt in VIEW_TYPES:
            img = get_image_for_car_type(master_car_type, vt)
            print(f"  {vt}: {img}")
    else:
        print("  Using default: side_view.png (no master type)")

    # ── Step 4: Generate PPT ──
    print("\n[4/5] Generating presentation...")
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Title slide
    print("  Creating title slide...")
    create_title_slide(prs, master_car, master_car_type, selected_cars,
                       type_groups, car_types, len(metric_names))

    # Metric slides
    for i, metric_name in enumerate(metric_names, 1):
        print(f"  Creating slide {i}/{len(metric_names)}: {metric_name}")
        create_metric_slide(
            prs, metric_name, master_car, master_car_type,
            selected_cars, all_cars_data, metric_config, arrow_config, i
        )

    # ── Step 5: Save ──
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    output_filename = f"Car_Comparison_{timestamp}.pptx"
    output_path = os.path.join(OUTPUT_DIR, output_filename)

    print(f"\n[5/5] Saving presentation...")
    prs.save(output_path)
    print(f"  Saved to: {output_path}")

    # Clean up temp files
    try:
        shutil.rmtree(TEMP_DIR)
        os.makedirs(TEMP_DIR, exist_ok=True)
    except Exception:
        pass

    total_cars = (1 if master_car else 0) + len(selected_cars)
    print("\n" + "=" * 60)
    print(f"  PPT generated successfully!")
    print(f"  File: {output_path}")
    print(f"  Slides: {len(metric_names) + 1} (1 title + {len(metric_names)} metrics)")
    if master_car:
        print(f"  Master car: {master_car} ({master_car_type})")
    print(f"  Comparison cars: {', '.join(selected_cars)}")
    print("=" * 60)
    input("\nPress Enter to exit...")


if __name__ == "__main__":
    main()
